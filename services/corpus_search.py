# -*- coding: utf-8 -*-
import hashlib
import os
import re
import sqlite3
from datetime import datetime
from pathlib import Path


DB_PATH = Path(__file__).resolve().parent.parent / "data" / "corpus_index.db"

_NLP = None
_NLP_MODE = "en_core_web_sm"


def _get_connection():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def _table_columns(conn, table_name):
    rows = conn.execute(f"PRAGMA table_info({table_name})").fetchall()
    return {str(row["name"]) for row in rows}


def ensure_schema():
    conn = _get_connection()
    try:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS documents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                path TEXT NOT NULL UNIQUE,
                name TEXT NOT NULL,
                file_type TEXT NOT NULL,
                source_type TEXT NOT NULL DEFAULT '',
                file_hash TEXT,
                imported_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS chunks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id INTEGER NOT NULL,
                chunk_text TEXT NOT NULL,
                page_num INTEGER,
                test_label TEXT,
                section_label TEXT,
                part_label TEXT,
                speaker_label TEXT,
                question_label TEXT,
                sort_key INTEGER NOT NULL,
                FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS sentences (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id INTEGER NOT NULL,
                chunk_id INTEGER,
                sentence_text TEXT NOT NULL,
                lemma_text TEXT NOT NULL,
                source_file TEXT NOT NULL,
                page_num INTEGER,
                test_label TEXT,
                section_label TEXT,
                part_label TEXT,
                speaker_label TEXT,
                question_label TEXT,
                sentence_order INTEGER NOT NULL DEFAULT 0,
                sort_key INTEGER NOT NULL,
                FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE,
                FOREIGN KEY(chunk_id) REFERENCES chunks(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS sentence_lemmas (
                sentence_id INTEGER NOT NULL,
                lemma TEXT NOT NULL,
                FOREIGN KEY(sentence_id) REFERENCES sentences(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS imports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id INTEGER,
                source_path TEXT NOT NULL,
                status TEXT NOT NULL,
                error_message TEXT,
                chunks_indexed INTEGER NOT NULL DEFAULT 0,
                sentences_indexed INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                finished_at TEXT,
                FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE SET NULL
            );
            """
        )

        document_columns = _table_columns(conn, "documents")
        if "source_type" not in document_columns:
            conn.execute("ALTER TABLE documents ADD COLUMN source_type TEXT NOT NULL DEFAULT ''")

        sentence_columns = _table_columns(conn, "sentences")
        if "chunk_id" not in sentence_columns:
            conn.execute("ALTER TABLE sentences ADD COLUMN chunk_id INTEGER")
        if "sentence_order" not in sentence_columns:
            conn.execute("ALTER TABLE sentences ADD COLUMN sentence_order INTEGER NOT NULL DEFAULT 0")

        conn.executescript(
            """
            CREATE INDEX IF NOT EXISTS idx_sentences_document_id ON sentences(document_id);
            CREATE INDEX IF NOT EXISTS idx_sentences_chunk_id ON sentences(chunk_id);
            CREATE INDEX IF NOT EXISTS idx_chunks_document_id ON chunks(document_id);
            CREATE INDEX IF NOT EXISTS idx_sentence_lemmas_lemma ON sentence_lemmas(lemma);
            CREATE UNIQUE INDEX IF NOT EXISTS idx_sentence_lemmas_unique ON sentence_lemmas(sentence_id, lemma);
            CREATE INDEX IF NOT EXISTS idx_imports_document_id ON imports(document_id);
            """
        )

        conn.commit()
    finally:
        conn.close()


def _file_hash(path):
    sha1 = hashlib.sha1()
    with open(path, "rb") as fp:
        while True:
            chunk = fp.read(1024 * 1024)
            if not chunk:
                break
            sha1.update(chunk)
    return sha1.hexdigest()


def get_nlp():
    global _NLP, _NLP_MODE
    if _NLP is not None:
        return _NLP, _NLP_MODE

    try:
        import spacy
    except Exception as e:
        raise RuntimeError(
            "spaCy is not installed. Run: pip install spacy spacy-lookups-data"
        ) from e

    try:
        _NLP = spacy.load("en_core_web_sm")
        return _NLP, _NLP_MODE
    except Exception as e:
        raise RuntimeError(
            "spaCy model 'en_core_web_sm' is missing. Run: python -m spacy download en_core_web_sm"
        ) from e


def get_nlp_status():
    _nlp, mode = get_nlp()
    return mode


def _clean_line(text):
    return re.sub(r"\s+", " ", str(text or "").strip())


def _looks_like_margin_label(text, x0, page_width):
    text = _clean_line(text)
    upper = text.upper()
    if not text or page_width in (None, 0):
        return False
    if x0 is None or float(x0) < float(page_width) * 0.72:
        return False
    return bool(
        re.fullmatch(r"Q\d{1,2}", upper)
        or upper == "EXAMPLE"
        or re.fullmatch(r"PAGE\s+\d+", upper)
    )


def _looks_like_speaker_label(text):
    return bool(re.fullmatch(r"([A-Z][A-Z '\-]{1,30}|[A-Za-z][A-Za-z '\-]{1,30}):", _clean_line(text)))


def _clean_pdf_block_text(text, x0, page_width):
    parts = []
    for raw_line in str(text or "").splitlines():
        line = _clean_line(raw_line)
        if not line:
            continue
        if _looks_like_margin_label(line, x0, page_width):
            continue
        if _looks_like_speaker_label(line):
            continue
        if re.fullmatch(r"\d{1,4}", line):
            continue
        parts.append(line)
    clean_text = _clean_line(" ".join(parts))
    clean_text = re.sub(r"\s+\bQ\d{1,2}\b\s*$", "", clean_text, flags=re.IGNORECASE)
    clean_text = re.sub(r"\s+\bEXAMPLE\b\s*$", "", clean_text, flags=re.IGNORECASE)
    return _clean_line(clean_text)


def _iter_docx_blocks(path):
    try:
        from docx import Document
        from docx.document import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except Exception as e:
        raise RuntimeError("python-docx is not installed. Run: pip install python-docx") from e

    document = Document(path)

    def _iter_block_items(parent):
        parent_elm = parent.element.body if isinstance(parent, DocxDocument) else parent._tc
        for child in parent_elm.iterchildren():
            if child.tag.endswith("}p"):
                yield Paragraph(child, parent)
            elif child.tag.endswith("}tbl"):
                yield Table(child, parent)

    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            text = _clean_line(block.text)
            if text:
                yield {"text": text, "page_num": None}
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [_clean_line(cell.text) for cell in row.cells]
                cells = [c for c in cells if c]
                if not cells:
                    continue
                if len(cells) == 1:
                    yield {"text": cells[0], "page_num": None}
                else:
                    yield {"text": "\t".join(cells[:2]), "page_num": None}


def _iter_pdf_blocks(path):
    try:
        import fitz
    except Exception as e:
        raise RuntimeError("PyMuPDF is not installed. Run: pip install pymupdf") from e

    with fitz.open(path) as doc:
        for page_index, page in enumerate(doc, start=1):
            page_width = float(page.rect.width or 0)
            text_data = page.get_text("dict")
            text_blocks = []
            for block in text_data.get("blocks", []):
                if block.get("type") != 0:
                    continue
                bbox = block.get("bbox") or (0, 0, 0, 0)
                x0, y0 = float(bbox[0] or 0), float(bbox[1] or 0)
                lines = []
                for line in block.get("lines", []):
                    spans = [span.get("text", "") for span in line.get("spans", [])]
                    line_text = _clean_line("".join(spans))
                    if line_text:
                        lines.append(line_text)
                clean_text = _clean_pdf_block_text("\n".join(lines), x0, page_width)
                if not clean_text:
                    continue
                text_blocks.append(
                    {
                        "text": clean_text,
                        "page_num": page_index,
                        "x0": x0,
                        "page_width": page_width,
                        "y0": y0,
                    }
                )
            text_blocks.sort(key=lambda item: (-item.get("y0", 0), item.get("x0", 0)))
            for item in text_blocks:
                yield item


def _iter_text_blocks(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as fp:
        for raw_line in fp:
            line = _clean_line(raw_line)
            if line:
                yield {"text": line, "page_num": None}


def _iter_file_blocks(path):
    lower = str(path).lower()
    if lower.endswith(".docx"):
        return list(_iter_docx_blocks(path))
    if lower.endswith(".pdf"):
        return list(_iter_pdf_blocks(path))
    return list(_iter_text_blocks(path))


def _is_continuation_line(raw_text, block, last_item, context):
    if not last_item:
        return False
    if block.get("page_num") != last_item.get("page_num"):
        return False
    if context["test_label"] != last_item.get("test_label"):
        return False
    if context["section_label"] != last_item.get("section_label"):
        return False
    if re.match(r"^([A-Z][A-Z '\-]{1,30}|[A-Za-z][A-Za-z '\-]{1,30}):\s*(.+)$", raw_text):
        return False
    if _looks_like_speaker_label(raw_text):
        return False

    prev_text = _clean_line(last_item.get("text"))
    if not prev_text:
        return False
    starts_lower = raw_text[:1].islower()
    prev_incomplete = prev_text[-1:] not in ".!?"
    if not (starts_lower or prev_incomplete):
        return False

    x0 = block.get("x0")
    page_width = block.get("page_width")
    last_x0 = last_item.get("_x0")
    if x0 is not None and page_width:
        if float(x0) < float(page_width) * 0.10:
            return False
        if last_x0 is not None and abs(float(x0) - float(last_x0)) > float(page_width) * 0.18 and not starts_lower:
            return False

    return True


def _parse_structured_blocks(blocks):
    parsed = []
    context = {
        "test_label": "",
        "section_label": "",
        "part_label": "",
        "speaker_label": "",
        "question_label": "",
    }
    last_item = None
    pending_speaker = ""
    for block in blocks:
        raw_text = _clean_line(block.get("text"))
        page_num = block.get("page_num")
        if not raw_text:
            continue

        upper = raw_text.upper()
        if re.fullmatch(r"TEST\s+\d+", upper):
            context["test_label"] = raw_text
            context["part_label"] = ""
            context["speaker_label"] = ""
            pending_speaker = ""
            continue
        if re.fullmatch(r"SECTION\s+\d+", upper) or re.fullmatch(r"PASSAGE\s+\d+", upper):
            context["section_label"] = raw_text
            context["speaker_label"] = ""
            pending_speaker = ""
            continue
        if len(raw_text) <= 40 and not re.search(r"[.!?]", raw_text) and raw_text == raw_text.title():
            context["part_label"] = raw_text
            continue

        question_match = re.search(r"\bQ\d+\b", raw_text, flags=re.IGNORECASE)
        if question_match:
            context["question_label"] = question_match.group(0).upper()
            raw_text = _clean_line(raw_text.replace(question_match.group(0), ""))
            if not raw_text:
                continue

        speaker_only_match = re.fullmatch(r"([A-Z][A-Z '\-]{1,30}|[A-Za-z][A-Za-z '\-]{1,30}):", raw_text)
        if speaker_only_match:
            pending_speaker = _clean_line(speaker_only_match.group(1))
            continue

        speaker_match = re.match(r"^([A-Z][A-Z '\-]{1,30}|[A-Za-z][A-Za-z '\-]{1,30}):\s*(.+)$", raw_text)
        if speaker_match:
            speaker = _clean_line(speaker_match.group(1))
            text = _clean_line(speaker_match.group(2))
            item = {
                "text": text,
                "page_num": page_num,
                "test_label": context["test_label"],
                "section_label": context["section_label"],
                "part_label": context["part_label"],
                "speaker_label": speaker,
                "question_label": context["question_label"],
                "_x0": block.get("x0"),
            }
            parsed.append(item)
            last_item = item
            pending_speaker = ""
            continue

        if _is_continuation_line(raw_text, block, last_item, context):
            last_item["text"] = _clean_line(f"{last_item['text']} {raw_text}")
            continue

        item = {
            "text": raw_text,
            "page_num": page_num,
            "test_label": context["test_label"],
            "section_label": context["section_label"],
            "part_label": context["part_label"],
            "speaker_label": pending_speaker,
            "question_label": context["question_label"],
            "_x0": block.get("x0"),
        }
        parsed.append(item)
        last_item = item
        pending_speaker = ""
    return parsed


def _doc_sentences(text):
    nlp, _mode = get_nlp()
    doc = nlp(text)
    raw_sents = [sent.text.strip() for sent in doc.sents if sent.text and sent.text.strip()]
    if raw_sents:
        merged = []
        for sent in raw_sents:
            if merged:
                prev = merged[-1]
                if sent[:1].islower() or prev[-1:] not in ".!?":
                    merged[-1] = _clean_line(f"{prev} {sent}")
                    continue
            merged.append(sent)
        return merged
    return [text.strip()] if str(text or "").strip() else []


def _lemma_doc(text):
    nlp, _mode = get_nlp()
    doc = nlp(text)
    lemmas = []
    for token in doc:
        if token.is_space or token.is_punct:
            continue
        lemma = str(getattr(token, "lemma_", "") or "").strip().lower()
        if lemma in ("-pron-", ""):
            lemma = str(token.text or "").strip().lower()
        lemmas.append(lemma)
    return lemmas


def _highlight_ranges(text, query, lemmas):
    clean_text = str(text or "")
    clean_query = _clean_line(query)
    target_lemmas = {str(lemma or "").strip().lower() for lemma in (lemmas or []) if str(lemma or "").strip()}
    ranges = []

    if clean_query:
        for match in re.finditer(re.escape(clean_query), clean_text, flags=re.IGNORECASE):
            ranges.append((match.start(), match.end()))

    if target_lemmas:
        nlp, _mode = get_nlp()
        doc = nlp(clean_text)
        for token in doc:
            if token.is_space or token.is_punct:
                continue
            lemma = str(getattr(token, "lemma_", "") or "").strip().lower()
            if lemma in ("-pron-", ""):
                lemma = str(token.text or "").strip().lower()
            if lemma in target_lemmas:
                ranges.append((int(token.idx), int(token.idx) + len(token.text)))

    if not ranges:
        return []

    ranges.sort(key=lambda item: (item[0], item[1]))
    merged = [list(ranges[0])]
    for start, end in ranges[1:]:
        last = merged[-1]
        if start <= last[1]:
            last[1] = max(last[1], end)
        else:
            merged.append([start, end])
    return [(start, end) for start, end in merged]


def _replace_document(conn, path, name, file_type, file_hash):
    row = conn.execute("SELECT id FROM documents WHERE path = ?", (path,)).fetchone()
    if row:
        document_id = int(row["id"])
        conn.execute(
            "DELETE FROM sentence_lemmas WHERE sentence_id IN (SELECT id FROM sentences WHERE document_id = ?)",
            (document_id,),
        )
        conn.execute("DELETE FROM sentences WHERE document_id = ?", (document_id,))
        conn.execute("DELETE FROM chunks WHERE document_id = ?", (document_id,))
        conn.execute(
            "UPDATE documents SET name = ?, file_type = ?, source_type = ?, file_hash = ?, imported_at = ? WHERE id = ?",
            (
                name,
                file_type,
                _infer_source_type(name, file_type),
                file_hash,
                datetime.now().isoformat(timespec="seconds"),
                document_id,
            ),
        )
        return document_id

    cur = conn.execute(
        "INSERT INTO documents(path, name, file_type, source_type, file_hash, imported_at) VALUES(?,?,?,?,?,?)",
        (
            path,
            name,
            file_type,
            _infer_source_type(name, file_type),
            file_hash,
            datetime.now().isoformat(timespec="seconds"),
        ),
    )
    return int(cur.lastrowid)


def _infer_source_type(name, file_type):
    label = f"{str(name or '').lower()} {str(file_type or '').lower()}"
    if "audio" in label or "transcript" in label or "script" in label:
        return "ielts_transcript"
    if "reading" in label or "passage" in label:
        return "ielts_reading"
    return "generic"


def _create_import_record(conn, path, document_id):
    cur = conn.execute(
        """
        INSERT INTO imports(document_id, source_path, status, created_at)
        VALUES(?, ?, ?, ?)
        """,
        (
            document_id,
            path,
            "running",
            datetime.now().isoformat(timespec="seconds"),
        ),
    )
    return int(cur.lastrowid)


def _finish_import_record(conn, import_id, status, chunks_indexed=0, sentences_indexed=0, error_message=""):
    conn.execute(
        """
        UPDATE imports
        SET status = ?, error_message = ?, chunks_indexed = ?, sentences_indexed = ?, finished_at = ?
        WHERE id = ?
        """,
        (
            status,
            error_message,
            int(chunks_indexed or 0),
            int(sentences_indexed or 0),
            datetime.now().isoformat(timespec="seconds"),
            int(import_id),
        ),
    )


def import_corpus_files(paths):
    ensure_schema()
    conn = _get_connection()
    summary = {
        "files": 0,
        "chunks": 0,
        "sentences": 0,
        "errors": [],
        "nlp_mode": get_nlp_status(),
    }
    try:
        for input_path in paths:
            path = os.path.abspath(str(input_path or "").strip())
            if not path or not os.path.exists(path):
                continue
            import_id = None
            try:
                file_hash = _file_hash(path)
                document_id = _replace_document(
                    conn,
                    path=path,
                    name=os.path.basename(path),
                    file_type=Path(path).suffix.lower().lstrip("."),
                    file_hash=file_hash,
                )
                import_id = _create_import_record(conn, path, document_id)
                raw_blocks = _iter_file_blocks(path)
                structured_blocks = _parse_structured_blocks(raw_blocks)
                sort_key = 0
                document_chunks = 0
                document_sentences = 0
                for block in structured_blocks:
                    chunk_text = _clean_line(block.get("text"))
                    if len(chunk_text) < 2:
                        continue
                    cur_chunk = conn.execute(
                        """
                        INSERT INTO chunks(
                            document_id, chunk_text, page_num, test_label, section_label,
                            part_label, speaker_label, question_label, sort_key
                        ) VALUES(?,?,?,?,?,?,?,?,?)
                        """,
                        (
                            document_id,
                            chunk_text,
                            block.get("page_num"),
                            block.get("test_label") or "",
                            block.get("section_label") or "",
                            block.get("part_label") or "",
                            block.get("speaker_label") or "",
                            block.get("question_label") or "",
                            sort_key,
                        ),
                    )
                    chunk_id = int(cur_chunk.lastrowid)
                    document_chunks += 1
                    sentence_order = 0
                    for sentence in _doc_sentences(chunk_text):
                        clean_sentence = _clean_line(sentence)
                        if len(clean_sentence) < 2:
                            continue
                        lemmas = _lemma_doc(clean_sentence)
                        lemma_text = " ".join(lemmas)
                        cur = conn.execute(
                            """
                            INSERT INTO sentences(
                                document_id, chunk_id, sentence_text, lemma_text, source_file, page_num,
                                test_label, section_label, part_label, speaker_label, question_label,
                                sentence_order, sort_key
                            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)
                            """,
                            (
                                document_id,
                                chunk_id,
                                clean_sentence,
                                lemma_text,
                                os.path.basename(path),
                                block.get("page_num"),
                                block.get("test_label") or "",
                                block.get("section_label") or "",
                                block.get("part_label") or "",
                                block.get("speaker_label") or "",
                                block.get("question_label") or "",
                                sentence_order,
                                sort_key,
                            ),
                        )
                        sentence_id = int(cur.lastrowid)
                        for lemma in sorted(set(lemmas)):
                            if not lemma:
                                continue
                            conn.execute(
                                "INSERT OR IGNORE INTO sentence_lemmas(sentence_id, lemma) VALUES(?, ?)",
                                (sentence_id, lemma),
                            )
                        sentence_order += 1
                        sort_key += 1
                        document_sentences += 1
                        summary["sentences"] += 1
                    summary["chunks"] += 1
                summary["files"] += 1
                _finish_import_record(
                    conn,
                    import_id,
                    "completed",
                    chunks_indexed=document_chunks,
                    sentences_indexed=document_sentences,
                )
            except Exception as e:
                summary["errors"].append(f"{os.path.basename(path)}: {e}")
                try:
                    if import_id:
                        _finish_import_record(conn, import_id, "failed", error_message=str(e))
                except Exception:
                    pass
        conn.commit()
        return summary
    finally:
        conn.close()


def corpus_stats():
    ensure_schema()
    conn = _get_connection()
    try:
        docs = conn.execute("SELECT COUNT(*) AS c FROM documents").fetchone()["c"]
        chunks = conn.execute("SELECT COUNT(*) AS c FROM chunks").fetchone()["c"]
        sentences = conn.execute("SELECT COUNT(*) AS c FROM sentences").fetchone()["c"]
        return {
            "documents": int(docs or 0),
            "chunks": int(chunks or 0),
            "sentences": int(sentences or 0),
            "nlp_mode": get_nlp_status(),
        }
    finally:
        conn.close()


def list_documents(limit=200):
    ensure_schema()
    conn = _get_connection()
    try:
        rows = conn.execute(
            """
            SELECT
                d.name,
                d.path,
                d.file_type,
                d.source_type,
                d.imported_at,
                COUNT(DISTINCT c.id) AS chunk_count,
                COUNT(DISTINCT s.id) AS sentence_count
            FROM documents d
            LEFT JOIN chunks c ON c.document_id = d.id
            LEFT JOIN sentences s ON s.document_id = d.id
            GROUP BY d.id
            ORDER BY d.imported_at DESC, d.name ASC
            LIMIT ?
            """,
            (int(limit),),
        ).fetchall()
        return [dict(row) for row in rows]
    finally:
        conn.close()


def search_corpus(query, limit=50, document_path=None):
    ensure_schema()
    q = _clean_line(query)
    if not q:
        return {"query": "", "lemmas": [], "results": [], "nlp_mode": get_nlp_status()}

    lemmas = [lemma for lemma in _lemma_doc(q) if lemma]
    if not lemmas:
        return {"query": q, "lemmas": [], "results": [], "nlp_mode": get_nlp_status()}

    conn = _get_connection()
    try:
        doc_path = _clean_line(document_path)
        if len(lemmas) == 1:
            sql = """
                SELECT
                    s.*,
                    c.chunk_text,
                    d.source_type,
                    d.path AS document_path,
                    d.imported_at
                FROM sentence_lemmas sl
                JOIN sentences s ON s.id = sl.sentence_id
                LEFT JOIN chunks c ON c.id = s.chunk_id
                LEFT JOIN documents d ON d.id = s.document_id
                WHERE sl.lemma = ?
            """
            params = [lemmas[0]]
            if doc_path:
                sql += " AND d.path = ?"
                params.append(doc_path)
            sql += """
                ORDER BY d.imported_at DESC, d.name ASC, s.page_num, s.sort_key, s.sentence_order
                LIMIT ?
            """
            params.append(int(limit))
            rows = conn.execute(sql, tuple(params)).fetchall()
        else:
            placeholders = ",".join("?" for _ in lemmas)
            sql = f"""
                WITH matched AS (
                    SELECT s.id
                    FROM sentences s
                    JOIN sentence_lemmas sl ON sl.sentence_id = s.id
                    LEFT JOIN documents d ON d.id = s.document_id
                    WHERE sl.lemma IN ({placeholders})
            """
            params = list(lemmas)
            if doc_path:
                sql += " AND d.path = ?"
                params.append(doc_path)
            sql += """
                    GROUP BY s.id
                    HAVING COUNT(DISTINCT sl.lemma) >= ?
                )
                SELECT
                    s.*,
                    c.chunk_text,
                    d.source_type,
                    d.path AS document_path,
                    d.imported_at
                FROM matched m
                JOIN sentences s ON s.id = m.id
                LEFT JOIN chunks c ON c.id = s.chunk_id
                LEFT JOIN documents d ON d.id = s.document_id
                ORDER BY d.imported_at DESC, d.name ASC, s.page_num, s.sort_key, s.sentence_order
                LIMIT ?
            """
            params.extend([len(set(lemmas)), int(limit)])
            rows = conn.execute(sql, tuple(params)).fetchall()

        results = []
        for row in rows:
            item = dict(row)
            item["match_type"] = "lemma"
            item["highlight_ranges"] = _highlight_ranges(item.get("sentence_text") or "", q, lemmas)
            results.append(item)
        return {
            "query": q,
            "lemmas": lemmas,
            "results": results,
            "nlp_mode": get_nlp_status(),
            "document_path": doc_path,
        }
    finally:
        conn.close()
